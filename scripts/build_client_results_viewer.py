#!/usr/bin/env python3
"""Build the normalized JSON payload for the 2023 client-results viewer."""

from __future__ import annotations

import argparse
import json
from pathlib import Path
from typing import Any

import yaml
from openpyxl import load_workbook
from pydantic import BaseModel, ConfigDict, Field

from leaders_db.research.registry import list_question_specs

PROJECT_ROOT = Path(__file__).resolve().parents[1]
DEFAULT_CONFIG = PROJECT_ROOT / "configs/client-viewers/2023-top20-v1.yaml"
CHAPTER_TITLES = {
    "1B": "Nuclear and existential responsibility",
    "2B": "International peace",
    "3B": "Domestic safety",
    "4B": "Political freedom",
    "5B": "Economic well-being",
    "6B": "Social well-being",
    "7B": "Integrity",
    "8B": "Effectiveness",
}


class ViewerConfig(BaseModel):
    """Validated source and output paths for one static viewer release."""

    model_config = ConfigDict(extra="forbid")

    schema_version: str
    release_id: str
    target_year: int
    batch_manifest: Path
    original_dossier_root: Path
    original_judge_root: Path
    repair_dossier_root: Path
    repair_judge_root: Path
    client_workbook: Path
    client_sheet: str
    client_country_column: int = Field(ge=1)
    client_score_columns: dict[str, int]
    client_country_names: dict[str, str]
    replacement_iso3: tuple[str, ...]
    corrected_judgments: tuple[Path, ...] = ()
    output_json: Path


def _rooted(path: Path) -> Path:
    return path if path.is_absolute() else PROJECT_ROOT / path


def load_config(path: Path) -> ViewerConfig:
    """Load and validate one viewer configuration file."""

    payload = yaml.safe_load(path.read_text(encoding="utf-8"))
    return ViewerConfig.model_validate(payload)


def _load_json(path: Path) -> dict[str, Any]:
    payload = json.loads(path.read_text(encoding="utf-8"))
    if not isinstance(payload, dict):
        raise ValueError(f"expected JSON object: {path}")
    return payload


def _find_dossiers(root: Path) -> dict[str, tuple[Path, dict[str, Any]]]:
    found: dict[str, tuple[Path, dict[str, Any]]] = {}
    for path in sorted(root.glob("jobs/*/attempts/*/dossier.json")):
        payload = _load_json(path)
        iso3 = str(payload.get("iso3", ""))
        if not iso3:
            continue
        found[iso3] = (path, payload)
    return found


def _find_judgments(root: Path) -> dict[str, dict[str, dict[str, Any]]]:
    results: dict[str, dict[str, dict[str, Any]]] = {}
    paths = sorted(root.glob("jobs/*/attempts/*/chapter-judgment.json"))
    for path in paths:
        payload = _load_json(path)
        chapter_id = str(payload["chapter_id"])
        results[chapter_id] = {
            str(item["iso3"]): item for item in payload.get("evaluations", [])
        }
    return results


def _load_corrected_judgments(
    paths: tuple[Path, ...],
) -> dict[str, dict[str, dict[str, Any]]]:
    output: dict[str, dict[str, dict[str, Any]]] = {}
    for path in paths:
        payload = _load_json(_rooted(path))
        for item in payload.get("evaluations", []):
            chapter_id = str(item["chapter_id"])
            iso3 = str(item["iso3"])
            output.setdefault(chapter_id, {})[iso3] = item
    return output


def _latest_failed_4b(root: Path) -> dict[str, dict[str, Any]]:
    paths = sorted(root.glob("jobs/*/attempts/*/chapter-judgment.pending.json"))
    candidates = [_load_json(path) for path in paths]
    candidates = [item for item in candidates if item.get("chapter_id") == "4B"]
    if not candidates:
        return {}
    evaluations = candidates[-1].get("evaluations", [])
    normalized: dict[str, dict[str, Any]] = {}
    for evaluation in evaluations:
        copy = dict(evaluation)
        confidence = copy.get("confidence_score")
        if isinstance(confidence, (int, float)) and 0 <= confidence <= 1:
            copy["confidence_score"] = round(float(confidence) * 100, 2)
            copy["viewer_normalization_note"] = (
                "Confidence was serialized as a fraction in a mechanically invalid "
                "candidate and normalized to the required 0-100 display scale."
            )
        normalized[str(copy["iso3"])] = copy
    return normalized


def _client_scores(config: ViewerConfig) -> dict[str, dict[str, float | None]]:
    workbook = load_workbook(
        _rooted(config.client_workbook), read_only=True, data_only=True
    )
    sheet = workbook[config.client_sheet]
    row_by_country: dict[str, tuple[Any, ...]] = {}
    for row in sheet.iter_rows(values_only=True):
        value = row[config.client_country_column - 1]
        if isinstance(value, str):
            row_by_country[value.strip()] = row
    output: dict[str, dict[str, float | None]] = {}
    for iso3, country_name in config.client_country_names.items():
        row = row_by_country.get(country_name)
        if row is None:
            raise ValueError(f"client workbook country not found: {iso3} / {country_name}")
        scores: dict[str, float | None] = {}
        for chapter_id, column in config.client_score_columns.items():
            raw = row[column - 1]
            scores[chapter_id] = float(raw) if isinstance(raw, (int, float)) else None
        output[iso3] = scores
    workbook.close()
    return output


def _question_texts() -> dict[str, str]:
    return {
        spec.methodology_id: spec.text
        for spec in list_question_specs()
        if len(spec.methodology_id) >= 4
        and spec.methodology_id[:2] in CHAPTER_TITLES
    }


def _score_appraisal(score: float) -> str:
    if score >= 8:
        return "strong"
    if score >= 6:
        return "generally positive, with important limitations"
    if score >= 4:
        return "mixed"
    return "poor"


def _referenced_claims(
    references: list[dict[str, Any]], evidence_by_id: dict[str, dict[str, Any]]
) -> list[str]:
    claims: list[str] = []
    for reference in references[:2]:
        evidence = evidence_by_id.get(str(reference.get("evidence_id", "")), {})
        claim = evidence.get("claim") or reference.get("explanation")
        if claim:
            claims.append(str(claim).rstrip("."))
    return claims


def _lower_sentence_opening(value: str) -> str:
    for opening in ("The ", "No ", "A ", "An "):
        if value.startswith(opening):
            return value[0].lower() + value[1:]
    if len(value) > 1 and value[0].isupper() and value[1].islower():
        return value[0].lower() + value[1:]
    return value


def _anchor_sentence(value: str, *, direction: str) -> str:
    """Render a judge's anchor explanation without duplicating its causal prefix."""

    text = value.strip().rstrip(".")
    if not text:
        return ""
    if text.lower().startswith("i did not "):
        return text[0].upper() + text[1:] + "."
    return f"I did not go {direction} because {_lower_sentence_opening(text)}."


def _reader_abstract(
    judgment: dict[str, Any],
    *,
    chapter_title: str,
    evidence_by_id: dict[str, dict[str, Any]],
) -> str:
    """Turn the saved analytical envelope into a self-contained public abstract."""

    ruler = str(judgment["ruler_name"])
    score = judgment.get("score_1_to_10")
    period = str(judgment["period_start_year"])
    if judgment["period_end_year"] != judgment["period_start_year"]:
        period += f"–{judgment['period_end_year']}"
    if score is None:
        reason = str(judgment.get("insufficient_evidence_reason") or "").rstrip(".")
        return (
            f"The available evidence does not support a responsible overall rating "
            f"of {ruler}'s {chapter_title.lower()} in {period}. {reason}. This is an "
            "evidence limitation, not a finding of either good or poor performance; "
            "the record remains unscored until stronger ruler-attributed evidence is "
            "available."
        )

    appraisal = _score_appraisal(float(score))
    paragraphs = [
        f"Overall, {ruler}'s record on {chapter_title.lower()} in {period} was "
        f"assessed as {appraisal}, resulting in a score of {score:g} out of 10."
    ]
    positives = _referenced_claims(
        list(judgment.get("decisive_positive_evidence", [])), evidence_by_id
    )
    negatives = _referenced_claims(
        list(judgment.get("decisive_negative_evidence", [])), evidence_by_id
    )
    if positives:
        paragraphs.append(
            "Evidence supporting a more favorable assessment includes the following: "
            + "; and ".join(positives)
            + "."
        )
    if negatives:
        paragraphs.append(
            "Evidence weighing against a stronger assessment includes the following: "
            + "; and ".join(negatives)
            + "."
        )
    attribution = str(judgment.get("ruler_attribution") or "").rstrip(".")
    if attribution:
        paragraphs.append(f"Responsibility and evidentiary limits matter here: {attribution}.")
    higher = str(judgment.get("higher_anchor_rejected") or "").rstrip(".")
    lower = str(judgment.get("lower_anchor_rejected") or "").rstrip(".")
    if higher and lower:
        paragraphs.append(
            f"Taken together, {_anchor_sentence(higher, direction='higher')} "
            f"{_anchor_sentence(lower, direction='lower')}"
        )
    elif higher or lower:
        reason = higher or lower
        direction = "higher" if higher else "lower"
        paragraphs.append(f"Taken together, {_anchor_sentence(reason, direction=direction)}")
    return "\n\n".join(paragraphs)


def _chapter_record(
    dossier: dict[str, Any],
    chapter_id: str,
    judgment: dict[str, Any] | None,
    questions: dict[str, str],
) -> dict[str, Any]:
    evidence_by_id = {
        str(item["evidence_id"]): item for item in dossier.get("evidence", [])
    }
    supplemental = list(judgment.get("supplemental_evidence", [])) if judgment else []
    evidence_by_id.update(
        {str(item["evidence_id"]): item for item in supplemental}
    )
    mappings: dict[str, set[str]] = {}
    for mapping in dossier.get("mappings", []):
        methodology_id = str(mapping.get("methodology_id", ""))
        if methodology_id.startswith(f"{chapter_id}."):
            mappings.setdefault(methodology_id, set()).add(str(mapping["evidence_id"]))
    lenses = []
    coverage_by_id = {
        str(item["methodology_id"]): item
        for item in dossier.get("coverage", [])
        if str(item.get("methodology_id", "")).startswith(f"{chapter_id}.")
    }
    for index in range(1, 11):
        methodology_id = f"{chapter_id}.{index}"
        coverage = coverage_by_id.get(methodology_id, {})
        evidence_ids = set(coverage.get("evidence_ids", [])) | mappings.get(
            methodology_id, set()
        )
        lenses.append(
            {
                "methodology_id": methodology_id,
                "question": questions.get(methodology_id, methodology_id),
                "status": coverage.get("status", "not_recorded"),
                "write_up": coverage.get("reason", "No lens-specific write-up recorded."),
                "evidence_ids": sorted(evidence_ids),
            }
        )
    chapter_evidence_ids = sorted(
        {evidence_id for lens in lenses for evidence_id in lens["evidence_ids"]}
    )
    chapter_evidence_ids.extend(
        evidence_id
        for evidence_id in (str(item["evidence_id"]) for item in supplemental)
        if evidence_id not in chapter_evidence_ids
    )
    public_judgment = dict(judgment) if judgment else None
    if public_judgment:
        public_judgment.pop("supplemental_evidence", None)
        public_judgment["reader_abstract"] = _reader_abstract(
            public_judgment,
            chapter_title=CHAPTER_TITLES[chapter_id],
            evidence_by_id=evidence_by_id,
        )
    recovery = None
    if public_judgment and public_judgment.get("score_1_to_10") is None:
        recoverable = (
            public_judgment.get("manual_review_reason_type") == "recoverable_null"
        )
        recovery = {
            "status": (
                "pending_corrected_flow_review"
                if recoverable
                else "substantive_review_required"
            ),
            "current_evidence_count": len(chapter_evidence_ids),
            "missing_or_weak_lenses": public_judgment.get(
                "missing_or_weak_lenses", []
            ),
            "requested_follow_up": public_judgment.get("manual_review_reason"),
            "maximum_targeted_rounds": 2 if recoverable else 0,
        }
    return {
        "chapter_id": chapter_id,
        "title": CHAPTER_TITLES[chapter_id],
        "lenses": lenses,
        "evidence": [
            evidence_by_id[evidence_id]
            for evidence_id in chapter_evidence_ids
            if evidence_id in evidence_by_id
        ],
        "judge": public_judgment,
        "recovery": recovery,
    }


def _mse(
    automated: dict[str, float | None], client: dict[str, float | None]
) -> tuple[float | None, int]:
    squared = [
        (automated[chapter] - client[chapter]) ** 2
        for chapter in CHAPTER_TITLES
        if automated.get(chapter) is not None and client.get(chapter) is not None
    ]
    return (round(sum(squared) / len(squared), 3), len(squared)) if squared else (None, 0)


def build_payload(config: ViewerConfig) -> dict[str, Any]:
    """Join client scores, dossiers, lenses, citations, and chapter judgments."""

    batch = yaml.safe_load(_rooted(config.batch_manifest).read_text(encoding="utf-8"))
    original_dossiers = _find_dossiers(_rooted(config.original_dossier_root))
    repair_dossiers = _find_dossiers(_rooted(config.repair_dossier_root))
    original_judgments = _find_judgments(_rooted(config.original_judge_root))
    repair_judgments = _find_judgments(_rooted(config.repair_judge_root))
    repair_judgments["4B"] = _latest_failed_4b(_rooted(config.repair_judge_root))
    corrected_judgments = _load_corrected_judgments(config.corrected_judgments)
    for chapter_id, evaluations in corrected_judgments.items():
        original_judgments.setdefault(chapter_id, {}).update(evaluations)
        repair_judgments.setdefault(chapter_id, {}).update(evaluations)
    client_by_iso3 = _client_scores(config)
    questions = _question_texts()
    rulers = []
    overall_squared: list[float] = []
    for case in batch["cases"]:
        iso3 = str(case["iso3"])
        use_repair = iso3 in config.replacement_iso3
        path, dossier = (repair_dossiers if use_repair else original_dossiers)[iso3]
        judgments = repair_judgments if use_repair else original_judgments
        automated = {
            chapter_id: (
                judgments.get(chapter_id, {}).get(iso3, {}).get("score_1_to_10")
            )
            for chapter_id in CHAPTER_TITLES
        }
        client = client_by_iso3[iso3]
        mse, comparable = _mse(automated, client)
        overall_squared.extend(
            (automated[chapter] - client[chapter]) ** 2
            for chapter in CHAPTER_TITLES
            if automated.get(chapter) is not None and client.get(chapter) is not None
        )
        rulers.append(
            {
                "iso3": iso3,
                "country_name": dossier["country_name"],
                "ruler_name": dossier["ruler_name"],
                "ruler_year_id": dossier["ruler_year_id"],
                "source_run": "repair" if use_repair else "original",
                "dossier_path": path.relative_to(PROJECT_ROOT).as_posix(),
                "automated_scores": automated,
                "client_scores": client,
                "mse": mse,
                "comparable_chapters": comparable,
                "chapters": [
                    _chapter_record(
                        dossier,
                        chapter_id,
                        judgments.get(chapter_id, {}).get(iso3),
                        questions,
                    )
                    for chapter_id in CHAPTER_TITLES
                ],
            }
        )
    return {
        "schema_version": "leaders_db_client_viewer_data_v1",
        "release_id": config.release_id,
        "target_year": config.target_year,
        "chapter_titles": CHAPTER_TITLES,
        "overall_mse": round(sum(overall_squared) / len(overall_squared), 3),
        "overall_comparable_cells": len(overall_squared),
        "rulers": rulers,
        "attributions_path": "docs/sources/attributions.md",
        "method_note": (
            "MSE is the mean of squared automated-minus-client score differences "
            "over cells where both scores are numeric. Lower is better; nulls are "
            "excluded, never treated as zero."
        ),
        "methodology_update": (
            "This provisional release predates the corrected evidence-preservation "
            "and attribution rules. Existing scores and citations remain frozen. "
            "Unscored chapters are flagged for review through the corrected flow: "
            "accepted evidence must survive formatting with its chapter routing, "
            "and formal governing responsibility may establish attribution outside "
            "the personal-integrity chapter."
        ),
    }


def parse_args() -> argparse.Namespace:
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("--config", type=Path, default=DEFAULT_CONFIG)
    return parser.parse_args()


def main() -> None:
    args = parse_args()
    config = load_config(args.config)
    payload = build_payload(config)
    output = _rooted(config.output_json)
    output.parent.mkdir(parents=True, exist_ok=True)
    output.write_text(json.dumps(payload, indent=2, ensure_ascii=False) + "\n", encoding="utf-8")
    print(f"wrote {output.relative_to(PROJECT_ROOT)} ({len(payload['rulers'])} rulers)")


if __name__ == "__main__":
    main()
